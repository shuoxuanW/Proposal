"""附录页渲染：在模板风格下动态新增检测效果/验证数据/风险/设备参数页。"""
from __future__ import annotations

from pathlib import Path

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.util import Inches, Pt

from .planner import SlideSpec
from .profile import TemplateProfile


def _rgb(value: str, default: str = "#1A5CD6") -> RGBColor:
    s = (value or default).lstrip("#")
    if len(s) != 6:
        s = default.lstrip("#")
    return RGBColor(int(s[0:2], 16), int(s[2:4], 16), int(s[4:6], 16))


def _move_last_slide(prs, index: int) -> None:
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    xml_slides.remove(slides[-1])
    xml_slides.insert(max(0, index), slides[-1])


def _clear(slide) -> None:
    for sh in list(slide.shapes):
        sh._element.getparent().remove(sh._element)


def _header(slide, title: str, style: dict) -> None:
    primary = _rgb(style.get("primary", "#1A5CD6"))
    font = style.get("font", "微软雅黑")
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.33), Inches(0.24),
                                 Inches(0.17), Inches(0.58))
    bar.fill.solid(); bar.fill.fore_color.rgb = primary
    bar.line.fill.background(); bar.shadow.inherit = False
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.12), Inches(8.0), Inches(0.59))
    run = tb.text_frame.paragraphs[0].add_run(); run.text = title
    run.font.size = Pt(int(style.get("title_size", 32)))
    run.font.name = font
    run.font.bold = True
    run.font.color.rgb = primary
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.42), Inches(0.82),
                                    Inches(13.2), Inches(0.82))
    ln.line.color.rgb = primary; ln.line.width = Pt(1.5)
    foot = slide.shapes.add_textbox(Inches(10.35), Inches(0.07), Inches(1.86), Inches(0.64))
    fr = foot.text_frame.paragraphs[0].add_run()
    fr.text = style.get("footer", "YD Vision  |  For Better Service")
    fr.font.size = Pt(9); fr.font.name = font; fr.font.color.rgb = _rgb("#999999")


def _textbox(slide, left, top, width, height, text: str, style: dict,
             size: int | None = None, color: str | None = None, bold: bool = False):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame; tf.word_wrap = True
    font = style.get("font", "微软雅黑")
    for i, line in enumerate(str(text or "").splitlines()):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = line
        r.font.size = Pt(int(size or style.get("body_size", 15)))
        r.font.name = font
        r.font.bold = bold
        r.font.color.rgb = _rgb(color or style.get("text", "#202020"))
        p.space_after = Pt(4)
    return tb


def _bullets(slide, left, top, width, height, bullets: list[str], style: dict, size: int = 14):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame; tf.word_wrap = True
    font = style.get("font", "微软雅黑")
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = "• " + str(b)
        r.font.size = Pt(size); r.font.name = font
        r.font.color.rgb = _rgb(style.get("text", "#202020"))
        p.space_after = Pt(7)
    return tb


def _table(slide, rows: list[list[str]], left: float, top: float, width: float, height: float,
           style: dict) -> None:
    if not rows:
        return
    nrows, ncols = len(rows), max(len(r) for r in rows)
    shape = slide.shapes.add_table(nrows, ncols, Inches(left), Inches(top),
                                   Inches(width), Inches(height))
    table = shape.table
    font = style.get("font", "微软雅黑")
    primary = _rgb(style.get("primary", "#1A5CD6"))
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            cell = table.cell(ri, ci)
            cell.text = str(row[ci]) if ci < len(row) else ""
            cell.margin_left = Inches(0.06); cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.03); cell.margin_bottom = Inches(0.03)
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.name = font
                    r.font.size = Pt(12 if ri else 13)
                    r.font.bold = ri == 0
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if ri == 0 else _rgb(style.get("text", "#202020"))
            if ri == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = primary


def _add_images(slide, images: list[str], left: float, top: float, max_w: float, max_h: float) -> None:
    if not images:
        return
    try:
        from PIL import Image
    except Exception:
        Image = None
    count = min(len(images), 2)
    slot_w = max_w / count
    for i, img_path in enumerate(images[:count]):
        w = slot_w - 0.12
        h = max_h
        if Image is not None:
            try:
                with Image.open(img_path) as im:
                    iw, ih = im.size
                ratio = min(w / iw, h / ih)
                w, h = iw * ratio, ih * ratio
            except Exception:
                pass
        x = left + i * slot_w + (slot_w - w) / 2
        y = top + (max_h - h) / 2
        slide.shapes.add_picture(img_path, Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def add_appendix_slides(prs, specs: list[SlideSpec], profile: TemplateProfile,
                        before_index: int) -> None:
    style = profile.style or {}
    layout_index = int(style.get("layout_index", 0))
    for i, spec in enumerate(specs):
        slide = prs.slides.add_slide(prs.slide_layouts[layout_index])
        _clear(slide)
        _header(slide, spec.title, style)
        idx = before_index + i
        _move_last_slide(prs, idx)

        has_images = bool(spec.images)
        body_width = 6.6 if has_images else float(style.get("content_width", 12.2))
        body_left = float(style.get("margin_left", 0.5))
        body_top = float(style.get("margin_top", 1.3))
        body_height = float(style.get("content_height", 5.5))

        if spec.table:
            table_h = min(4.8, 0.55 * len(spec.table) + 0.2)
            _table(slide, spec.table, body_left, body_top, body_width, table_h, style)
            if spec.body:
                _textbox(slide, body_left, body_top + table_h + 0.15, body_width,
                         max(0.6, body_height - table_h - 0.15), spec.body, style, size=12)
        else:
            y = body_top
            if spec.body:
                _textbox(slide, body_left, y, body_width, 2.5 if spec.bullets else body_height,
                         spec.body, style)
                y += 2.6 if spec.bullets else 0
            if spec.bullets:
                _bullets(slide, body_left, y, body_width, max(1.0, body_height - (y - body_top)),
                         spec.bullets, style)

        if has_images:
            _add_images(slide, spec.images, body_left + body_width + 0.2, body_top,
                        max(1.0, 13.1 - (body_left + body_width + 0.2)), 4.9)
