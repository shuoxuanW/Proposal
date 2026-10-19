"""FastAPI 服务：生成方案书 / RAG 查询 / 素材回流 / 导出下载 / 光路布局示意图。

启动：uvicorn src.mv_proposal.server:app --host 0.0.0.0 --port 8000
"""
from __future__ import annotations

import json
import re
import threading
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, JSONResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles

class _NoCacheStaticFiles(StaticFiles):
    async def get_response(self, path, scope):
        resp = await super().get_response(path, scope)
        resp.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
        return resp
from pydantic import BaseModel

app = FastAPI(title="机器视觉解决方案书自动生成系统", version="1.0.0")

# 静态资源：web 前端 + 产物目录
from .config import Config as _Config
_cfg = _Config.load()
_web_dir = _cfg.project_root / "web"
_output_dir = _cfg.output_dir
_output_dir.mkdir(parents=True, exist_ok=True)
if _web_dir.exists():
    app.mount("/web", _NoCacheStaticFiles(directory=str(_web_dir), html=True), name="web")
app.mount("/output", StaticFiles(directory=str(_output_dir)), name="output")
_station_dir = _cfg.project_root / "data" / "station_diagrams"
if _station_dir.exists():
    app.mount("/station-lib", _NoCacheStaticFiles(directory=str(_station_dir), html=False),
              name="station-lib")

_state: dict = {}
_lock = threading.Lock()


def get_context():
    """懒加载 + 常驻：模型与向量库只构建一次。"""
    with _lock:
        if "ctx" not in _state:
            from .config import Config
            from .context import build_context
            cfg = Config.load()
            _state["cfg"] = cfg
            _state["ctx"] = build_context(cfg)
        return _state["cfg"], _state["ctx"]


class GenerateReq(BaseModel):
    text: str
    llm: str | None = None


class OpticReq(BaseModel):
    text: str


class QueryReq(BaseModel):
    query: str
    top_k: int = 5
    collection: str = "proposals"


class SelectReq(BaseModel):
    text: str
    llm: str | None = None
    threshold: float | None = None


class ChatReq(BaseModel):
    text: str
    llm: str | None = None
    history: list[dict] = []


class StationQueryReq(BaseModel):
    query: str
    top_k: int = 12


@app.get("/health")
def health():
    return {"status": "ok", "name": app.title}


@app.get("/kb/stats")
def kb_stats():
    cfg, ctx = get_context()
    return {
        "embedding_provider": ctx.embedder.provider,
        "embedding_fallback": ctx.embedder.fallback_reason,
        "llm_provider": ctx.llm.provider,
        "images": ctx.image_lib.count(),
        "image_types": {t: sum(1 for e in ctx.image_lib.entries if e.device_type == t)
                        for t in ctx.image_lib.list_types()},
        "text_chunks": ctx.kb.count(),
        "collections": ctx.kb.collection_counts(),
        "stations": ctx.station_lib.count(),
        "station_projects": len({
            e.get("project", "") for e in ctx.station_lib.entries
            if e.get("project") and not re.fullmatch(r"slide\d+", e.get("project", ""))
        }),
        "vectorstore": str(cfg.vectorstore_path),
    }


def _ppt_outline(run_dir) -> list[dict]:
    """读取生成的 proposal.pptx，返回每页的标题/摘要/是否含图，供前端预览。"""
    ppt = Path(run_dir) / "proposal.pptx"
    if not ppt.exists():
        return []
    try:
        from pptx import Presentation
        prs = Presentation(str(ppt))
    except Exception:
        return []
    out = []
    for i, slide in enumerate(prs.slides, 1):
        texts = [sh.text_frame.text.strip() for sh in slide.shapes
                 if sh.has_text_frame and sh.text_frame.text.strip()]
        title = next((t for t in texts if len(t) <= 24), "")
        snippet = max(texts, key=len)[:90] if texts else ""
        has_image = any(sh.shape_type == 13 for sh in slide.shapes)
        out.append({"index": i, "title": title or f"第 {i} 页", "text": snippet, "has_image": has_image})
    return out


@app.post("/generate")
def generate(req: GenerateReq):
    cfg, ctx = get_context()
    from .pipeline import generate as run
    res = run(req.text, cfg=cfg, llm_provider=req.llm, out_dir=None, ctx=ctx)
    m = dict(res["manifest"])
    m["run_id"] = res["run_dir"].name
    m["run_dir"] = str(res["run_dir"])
    m["issues"] = [i.model_dump() for i in res["report"].issues]
    m["ppt_slides"] = _ppt_outline(res["run_dir"])
    # 历史工位图命中时补一个可直接预览的 URL（前端结果区显示缩略图）
    lr = m.get("layout_reasoning") or {}
    sh = lr.get("station_hit")
    if isinstance(sh, dict) and sh.get("id"):
        sh["image_url"] = "/station-lib/" + quote(str(sh["id"])) + ".png"
    return m


@app.post("/kb/query")
def kb_query(req: QueryReq):
    _, ctx = get_context()
    results = ctx.kb.retrieve(req.collection, req.query, top_k=req.top_k)
    return {"results": results}


def _station_item(e: dict) -> dict:
    return {
        "id": e["id"],
        "project": e.get("project", ""),
        "drawing_type": (e.get("data") or {}).get("drawing_type", ""),
        "image_url": "/station-lib/" + quote(str(e["id"])) + ".png",
    }


@app.get("/debug/status")
def debug_status():
    """给运维/调试用的只读状态快照。"""
    cfg, ctx = get_context()
    station_dir = cfg.project_root / "data" / "station_diagrams"
    projects = {
        e.get("project", "") for e in ctx.station_lib.entries
        if e.get("project") and not re.fullmatch(r"slide\d+", e.get("project", ""))
    }
    mtime = max((f.stat().st_mtime for f in station_dir.iterdir() if f.is_file()), default=None) if station_dir.exists() else None
    return {
        "version": app.version,
        "llm_provider": ctx.llm.provider,
        "llm_model": cfg.llm_model,
        "stations": ctx.station_lib.count(),
        "station_projects": len(projects),
        "station_library_mtime": datetime.fromtimestamp(mtime).isoformat() if mtime else None,
        "ppt_template": str(cfg.ppt_template),
        "ppt_profile": str(cfg.ppt_profile),
        "tmp_render_exists": (cfg.project_root / "_tmp_render").exists(),
        "collections": ctx.kb.collection_counts(),
    }


@app.get("/station-library/list")
def station_library_list(limit: int = 200, offset: int = 0):
    """列出历史工位图库，供前端图库面板浏览。"""
    _, ctx = get_context()
    entries = sorted(ctx.station_lib.entries, key=lambda e: e["id"])
    page = entries[max(0, offset):max(0, offset) + max(1, min(limit, 500))]
    return {"total": len(entries), "items": [_station_item(e) for e in page]}


@app.post("/station-library/search")
def station_library_search(req: StationQueryReq):
    """按需求关键词检索历史工位图。"""
    _, ctx = get_context()
    hits = ctx.station_lib.search(req.query, top_k=max(1, min(req.top_k, 50)))
    return {
        "total": len(hits),
        "items": [
            {**_station_item(h), "score": round(h["score"], 4)}
            for h in hits
        ],
    }


def _select_ctx(ctx, llm_provider):
    """选型请求可指定 LLM：复用缓存的知识库，仅重建轻量 LLM 客户端（不重载向量模型）。"""
    if not llm_provider:
        return ctx
    from types import SimpleNamespace
    from .llm import LLMClient
    cfg = _state["cfg"]
    llm = LLMClient(provider=llm_provider, base_url=cfg.llm_base_url,
                    api_key=cfg.api_key, model=cfg.llm_model, temperature=cfg.llm_temperature)
    return SimpleNamespace(kb=ctx.kb, llm=llm)


@app.post("/chat")
def chat(req: ChatReq):
    """售前选型顾问：客户侧只返回自然语言对话；检索详情在 debug 里单独返回。"""
    _, ctx = get_context()
    from .advisor import advise
    sctx = _select_ctx(ctx, req.llm)
    return advise(sctx, req.history or [], req.text, sctx.llm)


@app.post("/select")
def select(req: SelectReq):
    """选型：案例优先召回 → 动态设备组合兜底（含硬件强制配套校验）。"""
    _, ctx = get_context()
    from .selection import format_selection, run_selection
    sctx = _select_ctx(ctx, req.llm)
    kwargs = {} if req.threshold is None else {"threshold": req.threshold}
    res = run_selection(sctx, req.text, **kwargs)
    return {
        "source": res.source,
        "markdown": format_selection(res),
        "matched": [{"source": c["source"], "score": round(c["score"], 4)} for c in res.matched_cases],
    }


@app.post("/kb/search-images")
def search_images(query: str, device_type: str, top_k: int = 5):
    _, ctx = get_context()
    return {"results": ctx.image_lib.search(device_type, query, top_k=top_k)}


@app.post("/optic")
def optic(req: OpticReq):
    """光路/工位图：① 检索历史工位示意图（命中→重画）；② 否则 LLM 推理布局；③ 失败回退规则光路图。"""
    cfg, ctx = get_context()
    from . import schematic

    out = cfg.output_dir / "optic"
    out.mkdir(parents=True, exist_ok=True)

    # ① 检索历史工位示意图
    station = getattr(ctx, "station_lib", None)
    if station:
        hits = station.search(req.text, top_k=1)
        if hits and hits[0]["score"] >= 0.45:
            h = hits[0]
            png, svg = schematic.render_schematic(h["data"], out, name="optic_layout")
            return {"kind": "station_image", "matched": h["id"], "score": round(h["score"], 3),
                    "png": str(out / png), "svg": str(out / svg),
                    "layout_reasoning": {"status": "retrieved", "reason": "命中历史工位图", "alternatives": []}}

    # ② 降级：LLM 推理布局 / 规则光路图
    from . import diagram
    from .station import ALTERNATIVE_MODELS, reason_station_layout

    spec, ok, reason = reason_station_layout(req.text, ctx.llm)
    if ok and spec is not None:
        tl = getattr(ctx, "template_lib", None)
        if tl:
            for cam in spec.cameras:
                hit = tl.best("camera", cam.label)
                if hit:
                    cam.photo = hit["rel_path"]
            for lt in spec.lights:
                hit = tl.best("light_source", lt.label)
                if hit:
                    lt.photo = hit["rel_path"]
        png, svg = diagram.render_station_layout(
            spec, out, image_root=cfg.template_dir, dpi=cfg.image_dpi, name="station")
        return {"kind": "station", "camera_count": spec.camera_count,
                "cameras": [c.label for c in spec.cameras],
                "lights": [l.label for l in spec.lights],
                "png": str(out / png), "svg": str(out / svg),
                "layout_reasoning": {"status": "success", "reason": "", "alternatives": []}}

    from .optic import extract_optic_layout
    ospec = extract_optic_layout(req.text)
    png, svg = schematic.render_schematic(schematic.build_optic_json(ospec), out, name="optic_layout")
    return {"kind": "optic", "png": str(out / png), "svg": str(out / svg),
            "layout_reasoning": {"status": "mock" if ctx.llm.is_mock else "failed",
                                 "reason": reason, "alternatives": ALTERNATIVE_MODELS}}

@app.post("/feedback/image")
def feedback_image(file: UploadFile = File(...), device_type: str = Form("accessory"),
                   name: str = Form(""), description: str = Form(""), params: str = Form("{}"),
                   tags: str = Form("")):
    _, ctx = get_context()
    import tempfile
    meta = {
        "device_type": device_type,
        "name": name or Path(file.filename or "image").stem,
        "description": description,
        "params": json.loads(params or "{}"),
        "tags": [t for t in tags.split(",") if t],
    }
    with tempfile.NamedTemporaryFile(suffix=Path(file.filename or "img").suffix, delete=False) as tf:
        tf.write(file.file.read())
        tmp = Path(tf.name)
    try:
        entry = ctx.image_lib.register(tmp, meta)
    finally:
        tmp.unlink(missing_ok=True)
    return entry.to_dict()


@app.post("/feedback/text")
def feedback_text(payload: dict):
    _, ctx = get_context()
    coll = payload.get("collection", "proposals")
    chunk = ctx.kb.add(coll, payload.get("text", ""), payload.get("source", "人工修改"))
    return {"id": chunk.id, "source": chunk.source, "collection": chunk.collection}


@app.get("/outputs")
def list_outputs():
    cfg, _ = get_context()
    runs = []
    for d in sorted(cfg.output_dir.iterdir()):
        if d.is_dir() and (d / "manifest.json").exists():
            runs.append({"run_id": d.name, "dir": str(d)})
    return {"runs": runs}


@app.get("/")
def index():
    return RedirectResponse("/web/")