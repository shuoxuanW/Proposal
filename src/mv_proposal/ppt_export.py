"""按《视觉方案书格式.pptx》模板生成 PPT 方案书。

模板 8 页：封面 / 目录 / 版本审阅 / 客户需求 / 检测说明 / 打光示意图-相机1 / 检测说明-相机2 / Thanks。
"""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

from .recommend import recommend_devices

# 第 6 页要替换掉的模板示意图相关形状（换成我们生成的光路布局示意图）
_SLIDE6_SCHEMATIC_SHAPES = [
    "组合 6147", "组合 2", "直接连接符 109", "直接箭头连接符 110",
    "直接连接符 113", "直接箭头连接符 116", "直接连接符 115", "矩形 114",
    "文本框 117", "文本框 111", "文本框 112",
]

# 上一个项目的内容（产品图/检测示意/俯视图），生成新方案时需清除
_SLIDE4_REMOVE = ["对象 2", "文本框 1"]
_SLIDE5_REMOVE = [
    "组合 10", "矩形 14", "矩形 15", "文本框 16", "椭圆 17", "矩形 19",
    "椭圆 21", "椭圆 22", "椭圆 27", "椭圆 52", "椭圆 53", "椭圆 54",
    "矩形 55", "直接连接符 56", "直接连接符 61",
    "矩形 62", "矩形 63", "矩形 64", "矩形 65", "矩形 66", "矩形 67",
    "矩形 68", "矩形 69", "矩形 70", "矩形 71", "矩形 72", "矩形 73",
    "矩形 74", "矩形 75", "矩形 76", "矩形 77", "矩形 78", "矩形 79",
    "文本框 80",
]
_SLIDE7_REMOVE = [
    "椭圆 18", "矩形 20", "椭圆 23", "椭圆 24", "椭圆 25", "椭圆 28",
    "椭圆 29", "椭圆 30", "矩形 26", "直接连接符 31", "直接连接符 32",
    "矩形 33", "文本框 51", "椭圆 12", "椭圆 13",
]


def _find(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


def set_text(shape, text: str) -> None:
    """替换文本框文字，尽量保留原首段 run 的字体样式。"""
    tf = shape.text_frame
    src = None
    for p in tf.paragraphs:
        for r in p.runs:
            if r.text.strip():
                src = r
                break
        if src:
            break
    tf.clear()
    for i, line in enumerate(text.split("\n")):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = para.add_run()
        run.text = line
        if src is not None and src.font is not None:
            f = src.font
            run.font.name = f.name
            run.font.size = f.size
            run.font.bold = f.bold
            run.font.italic = f.italic
            try:
                if f.color and f.color.type is not None:
                    run.font.color.rgb = f.color.rgb
            except Exception:
                pass


def replace_runs(shape, mapping: dict[str, str]) -> None:
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            for k, v in mapping.items():
                if k in r.text:
                    r.text = r.text.replace(k, v)


def _delete_shapes(slide, names: list[str]) -> None:
    for sh in list(slide.shapes):
        if sh.name in names:
            sh._element.getparent().remove(sh._element)


def _optic_spec(plan):
    if plan.diagrams and plan.diagrams[0].kind == "optic" and getattr(plan.diagrams[0], "optic_spec", None):
        return plan.diagrams[0].optic_spec
    return None


def _artifact_image(artifacts, kind: str, out_dir: Path):
    for a in artifacts:
        if a.kind == kind:
            p = out_dir / a.png_rel
            if p.exists():
                return p
    return None


def _short_name(name: str) -> str:
    s = re.sub(r"(视觉)?方案书?$", "", name).strip()
    return s or name


def _strip_orphan_media(pptx_path: Path) -> None:
    """剔除未被引用的 OLE/WMF/VML 等孤立关系与部件，避免 WPS 渲染模板残留图。"""
    import re
    import zipfile

    pptx_path = Path(pptx_path)
    with zipfile.ZipFile(pptx_path) as z:
        names = z.namelist()
        data = {n: z.read(n) for n in names}
    new_data = dict(data)

    # 1) 每个 slide：删除 slide xml 未引用的 image/oleObject/vmlDrawing 关系
    for n in names:
        m = re.match(r"ppt/slides/slide(\d+)\.xml$", n)
        if not m:
            continue
        xml = data[n].decode("utf-8")
        rids = set(re.findall(r"r:(?:embed|id|link)=\"(rId\d+)\"", xml))
        rels_path = f"ppt/slides/_rels/slide{m.group(1)}.xml.rels"
        if rels_path not in data:
            continue
        rels = data[rels_path].decode("utf-8")
        keep = []
        for rel in re.findall(r"<Relationship [^>]*/>", rels):
            rid = re.search(r'Id="(rId\d+)"', rel).group(1)
            typ = re.search(r'Type="[^"]*?/(\w+)"', rel).group(1)
            if typ in ("image", "oleObject", "vmlDrawing") and rid not in rids:
                continue
            keep.append(rel)
        head = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        new_data[rels_path] = (head + "".join(keep) + "</Relationships>").encode("utf-8")

    # 2-3) 两轮清理：删除孤立部件 + 其附属 _rels（集合避免重复删除）
    for _ in range(2):
        referenced = set()
        for n, d in new_data.items():
            if n.endswith(".rels"):
                for t in re.findall(r'Target="\.\./([^"]+)"', d.decode("utf-8", "ignore")):
                    referenced.add(t)
        to_remove = set()
        for n in new_data:
            if n.startswith("ppt/media/") or n.startswith("ppt/embeddings/"):
                if n[len("ppt/"):] not in referenced:
                    to_remove.add(n)
            elif n.startswith("ppt/drawings/") and "/_rels/" not in n:
                if n[len("ppt/"):] not in referenced:
                    to_remove.add(n)
                    to_remove.add(f"ppt/drawings/_rels/{n.rsplit('/', 1)[1]}.rels")
        for n in new_data:
            if n.startswith("ppt/drawings/_rels/"):
                base = "ppt/drawings/" + n.rsplit("/", 1)[1][:-5]
                if base not in new_data or base in to_remove:
                    to_remove.add(n)
        for n in to_remove:
            if n in new_data:
                new_data.pop(n)

    tmp = pptx_path.with_suffix(".tmp.pptx")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in sorted(new_data):
            zout.writestr(n, new_data[n])
    pptx_path.unlink()
    tmp.rename(pptx_path)


def _add_selection_slide(prs, recs, index, profile=None) -> None:
    """新增「设备选型建议」页，插入到指定位置。"""
    style = profile.style if profile else {}
    layout_index = int((profile.base("selection").get("layout_index", 0) if profile else 0))
    slide = prs.slides.add_slide(prs.slide_layouts[layout_index])
    xml_slides = prs.slides._sldIdLst
    slds = list(xml_slides)
    xml_slides.remove(slds[-1])
    xml_slides.insert(index, slds[-1])
    for sh in list(slide.shapes):
        sh._element.getparent().remove(sh._element)
    from .ppt.appendix import _rgb
    primary = _rgb(style.get("primary", "#1A5CD6"))
    font = style.get("font", "微软雅黑")
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.33), Inches(0.24), Inches(0.17), Inches(0.58))
    bar.fill.solid(); bar.fill.fore_color.rgb = primary; bar.line.fill.background(); bar.shadow.inherit = False
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.12), Inches(6.0), Inches(0.59))
    run = tb.text_frame.paragraphs[0].add_run(); run.text = "设备选型建议"
    run.font.size = Pt(int(style.get("title_size", 32))); run.font.name = font; run.font.color.rgb = primary
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.42), Inches(0.82), Inches(13.2), Inches(0.82))
    ln.line.color.rgb = primary; ln.line.width = Pt(1.5)
    # 内容
    tb2 = slide.shapes.add_textbox(Inches(0.6), Inches(1.3), Inches(12.2), Inches(5.8))
    tf = tb2.text_frame; tf.word_wrap = True
    for i, r in enumerate(recs):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = para.add_run(); run.text = "• " + r
        run.font.size = Pt(16); run.font.name = font; run.font.color.rgb = _rgb(style.get("text", "#202020"))
        para.space_after = Pt(10)


def _blank_slide4(prs, index: int = 3) -> None:
    """客户未提供产品信息时，第4页留白（不虚构图纸、不套旧模板图）。"""
    slide = prs.slides[index]
    for sh in list(slide.shapes):
        sh._element.getparent().remove(sh._element)
    # 蓝色标题栏
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.33), Inches(0.24), Inches(0.17), Inches(0.58))
    bar.fill.solid(); bar.fill.fore_color.rgb = RGBColor(0x1A, 0x5C, 0xD6); bar.line.fill.background(); bar.shadow.inherit = False
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.12), Inches(5.6), Inches(0.59))
    run = tb.text_frame.paragraphs[0].add_run(); run.text = "客户需求"
    run.font.size = Pt(32); run.font.name = "微软雅黑"; run.font.color.rgb = RGBColor(0x1A, 0x5C, 0xD6)
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.42), Inches(0.82), Inches(13.2), Inches(0.82))
    ln.line.color.rgb = RGBColor(0x1A, 0x5C, 0xD6); ln.line.width = Pt(1.5)
    # 留白提示
    tb2 = slide.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(12.0), Inches(4.0))
    tf = tb2.text_frame; tf.word_wrap = True
    r = tf.paragraphs[0].add_run()
    r.text = "客户暂未提供产品相关信息，本页留白，待补充后自动生成产品示意图。"
    r.font.size = Pt(16); r.font.name = "微软雅黑"; r.font.color.rgb = RGBColor(0x99, 0x99, 0x99)


def _rebuild_slide4(prs, plan, wp_img, index: int = 3) -> None:
    """彻底重建第4页：删光所有形状（含 OLE/VML/WMF 残留），重画标题/内容/产品图。"""
    slide = prs.slides[index]
    for sh in list(slide.shapes):
        sh._element.getparent().remove(sh._element)
    # 蓝色竖条
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.33), Inches(0.24), Inches(0.17), Inches(0.58))
    bar.fill.solid(); bar.fill.fore_color.rgb = RGBColor(0x1A, 0x5C, 0xD6); bar.line.fill.background()
    bar.shadow.inherit = False
    # 标题
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.12), Inches(5.6), Inches(0.59))
    run = tb.text_frame.paragraphs[0].add_run(); run.text = "客户需求"
    run.font.size = Pt(32); run.font.name = "微软雅黑"; run.font.color.rgb = RGBColor(0x1A, 0x5C, 0xD6)
    # 分隔线
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.42), Inches(0.82), Inches(13.2), Inches(0.82))
    ln.line.color.rgb = RGBColor(0x1A, 0x5C, 0xD6); ln.line.width = Pt(1.5)
    # 内容
    reqs = plan.requirements or ["（待填写检测内容）"]
    content = (f"产品名称：{plan.product_name or '（待填写）'}\n"
               f"产品尺寸：{plan.product_size or '（待填写）'}\n"
               f"检测要求：\n" + "\n".join(f"{i + 1}. {r}" for i, r in enumerate(reqs)))
    tb2 = slide.shapes.add_textbox(Inches(0.5), Inches(1.29), Inches(7.0), Inches(6.1))
    tf = tb2.text_frame; tf.word_wrap = True
    for i, line in enumerate(content.split("\n")):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = para.add_run(); r.text = line
        r.font.size = Pt(15); r.font.name = "微软雅黑"; r.font.color.rgb = RGBColor(0x20, 0x20, 0x20)
    # 产品图
    if wp_img is not None:
        slide.shapes.add_picture(str(wp_img), Inches(7.55), Inches(1.4), width=Inches(5.3))


def export_ppt(plan, artifacts, template_path: Path, out_dir: Path, out_name: str = "proposal.pptx",
               chapters=None, report=None, profile_path=None) -> str:
    """动态生成 PPT：模板只提供风格；页序由 planner 决定。

    - 基础页：封面/目录/版本审阅/客户需求/检测说明/打光示意图/相机2/Thanks
    - 动态页：设备选型建议 + 附录（检测效果/验证数据/风险/设备参数）
    - 模板页序、形状名、样式全部来自 visual_proposal_profile.yaml
    """
    from .ppt.profile import load_profile
    from .ppt.planner import plan_appendix_slides
    from .ppt.appendix import add_appendix_slides

    profile = load_profile(profile_path)
    prs = Presentation(str(template_path))
    slides = list(prs.slides)
    date = datetime.now().strftime("%Y.%m.%d")
    ospec = _optic_spec(plan)

    def slide_for(role: str, default: int):
        idx = int(profile.base(role).get("index", default))
        return slides[idx] if 0 <= idx < len(slides) else slides[default]

    # ---- 封面 ----
    cover_cfg = profile.base("cover")
    s1 = slide_for("cover", 0)
    title = _find(s1, cover_cfg.get("title_shape", "标题 3"))
    if title is not None:
        replace_runs(title, {cover_cfg.get("title_token", "XXXX"): _short_name(plan.project_name)})
    sub = _find(s1, cover_cfg.get("subtitle_shape", "副标题 4"))
    if sub is not None:
        set_text(sub, f"客户名称：{plan.customer_name or '（此处填写完整的客户名称）'}\n"
                      f"技术窗口：{plan.tech_window or '（打光测试谁做的写谁）'}\n"
                      f"营业窗口：刘志民\n营业电话：18929323022\n汇报时间：{date}")

    # ---- 客户需求 ----
    req_idx = int(profile.base("requirement").get("index", 3))
    wp_img = _artifact_image(artifacts, "workpiece", out_dir)
    if wp_img is not None or plan.product_name or plan.requirements:
        _rebuild_slide4(prs, plan, wp_img, req_idx)
        if wp_img is not None:
            print(f"  [PPT] 客户需求页：已插入产品图 {wp_img.name}")
        else:
            print("  [PPT] 客户需求页：已填写产品/检测信息（无产品图）")
    else:
        _blank_slide4(prs, req_idx)
        print("  [PPT] 客户需求页：客户未提供产品信息，留白")

    # ---- 检测说明 ----
    det_cfg = profile.base("detection")
    s5 = slide_for("detection", 4)
    _delete_shapes(s5, det_cfg.get("remove_shapes", _SLIDE5_REMOVE))
    det_img = _artifact_image(artifacts, "detection", out_dir)
    if det_img is not None:
        s5.shapes.add_picture(str(det_img), Inches(9.55), Inches(1.45), width=Inches(3.3))
        print(f"  [PPT] 检测说明页：已插入检测示意图 {det_img.name}")
    box81 = _find(s5, det_cfg.get("body_shape", "文本框 81"))
    if box81 is not None:
        req_summary = "、".join(plan.requirements) or plan.detection_target
        parts = [f"综述：相机1对{plan.detection_target}进行检测，检测内容：{req_summary}。", ""]
        labels = ["检测内容一", "检测内容二", "检测内容三", "检测内容四"]
        for i, r in enumerate(plan.requirements):
            parts.append(f"{labels[i] if i < 4 else f'检测内容{i + 1}'}：{r}")
        parts.append("检测说明：系统采用「图像采集 → 图像处理 → 结果输出」三层架构。")
        if ospec is not None:
            parts.append(f"光路布局：{ospec.camera_label}（{ospec.camera_note}）垂直向下，成像光路对准{ospec.object_label}，"
                         f"镜头端面到检测物基准平面 {ospec.lens_to_object}；{ospec.light_color}{ospec.light_label}正视图摆放，"
                         f"角度 {int(ospec.tilt_min)}°~{int(ospec.tilt_max)}° 可调，底部距基准平台 {ospec.light_height}。")
        set_text(box81, "\n".join(parts))

    # ---- 打光/工位示意图 ----
    opt_cfg = profile.base("optical")
    s6 = slide_for("optical", 5)
    _delete_shapes(s6, opt_cfg.get("remove_shapes", _SLIDE6_SCHEMATIC_SHAPES))
    img_path = _artifact_image(artifacts, "station", out_dir) or _artifact_image(artifacts, "optic", out_dir)
    if img_path is not None:
        s6.shapes.add_picture(str(img_path), Inches(8.5), Inches(1.4), height=Inches(4.9))
        print(f"  [PPT] 打光示意图页：已插入工位架构图 {img_path.name}")
    table = _find(s6, opt_cfg.get("table_shape", "表格 124"))
    if table is not None and table.has_table:
        m = plan.key_metrics
        cc = getattr(plan, "case_constraints", None) or {}
        config = {
            "相机": cc.get("camera_model") or "工业相机（按案例/目录选型）",
            "镜头": cc.get("lens") or "镜头（按视场/工作距选型）",
            "光源": cc.get("light_type") or (f"{ospec.light_color}{ospec.light_label}" if ospec else "光源（待选型）"),
            "视野": m.get("视场") or cc.get("fov") or "（待填）",
            "分辨率": m.get("精度") or cc.get("accuracy") or "（待填）",
            "定位/检测精度": m.get("精度") or cc.get("accuracy") or "±0.1mm",
        }
        t = table.table
        for r in range(1, len(t.rows)):
            key = t.cell(r, 0).text.strip()
            if key in config:
                t.cell(r, 1).text = config[key]

    # ---- 相机2说明 ----
    cam2_cfg = profile.base("camera2")
    s7 = slide_for("camera2", 6)
    _delete_shapes(s7, cam2_cfg.get("remove_shapes", _SLIDE7_REMOVE))
    box3 = _find(s7, cam2_cfg.get("body_shape", "文本框 3"))
    if box3 is not None:
        set_text(box3, "检测内容：本项目为单相机方案（相机1），检测内容同前页。\n\n"
                       "检测说明：如后续扩展第二相机，再补充独立检测说明。")

    # ---- 版本审阅 ----
    ver_cfg = profile.base("version")
    s3 = slide_for("version", 2)
    tbl = _find(s3, ver_cfg.get("table_shape", "表格 2"))
    if tbl is not None and tbl.has_table:
        t = tbl.table
        row, col = int(ver_cfg.get("date_row", 1)), int(ver_cfg.get("date_col", 3))
        if len(t.rows) > row and len(t.columns) > col:
            t.cell(row, col).text = date

    # ---- 设备选型建议页 ----
    recs = recommend_devices(plan)
    sel_cfg = profile.base("selection")
    _add_selection_slide(prs, recs, int(sel_cfg.get("insert_after", 4)) + 1, profile)
    print(f"  [PPT] 新增「设备选型建议」页（{len(recs)} 条）")

    # ---- 动态附录页（模板风格） ----
    try:
        specs = plan_appendix_slides(plan, chapters or plan.chapters, artifacts, report, out_dir)
        if specs:
            add_appendix_slides(prs, specs, profile, len(prs.slides) - 1)
            print("  [PPT] 新增附录页：" + "、".join(s.title for s in specs))
    except Exception as e:
        print(f"  [PPT] 附录页生成失败：{e}")

    out_path = out_dir / out_name
    prs.save(str(out_path))
    _strip_orphan_media(out_path)
    return out_name
