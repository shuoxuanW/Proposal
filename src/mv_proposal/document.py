"""文档导出：Markdown / HTML / Word(.docx) / PDF。Word、PDF 接口预留并可离线生成。"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from .models import Proposal

IMG_RE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")
BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def compose_markdown(proposal: Proposal, resolved: list) -> str:
    parts = [
        f"# {proposal.project_name}",
        "",
        f"> 需求概要：{proposal.requirement_summary}",
        "",
    ]
    for ch in resolved:
        parts.append(ch.content.strip())
        parts.append("")
    parts.append("## 附图清单")
    parts.append("")
    for i, d in enumerate(proposal.diagrams, 1):
        parts.append(f"- 图{i}：{d.caption}（{d.png_rel}）")
    parts.append("")
    parts.append("## 校验审核报告")
    parts.append("")
    parts.append(f"校验结果：{'通过' if proposal.report.passed else '未通过'}，"
                 f"共 {len(proposal.report.issues)} 条问题。")
    for it in proposal.report.issues:
        parts.append(f"- [{it.severity}] {it.code}：{it.message}")
    return "\n".join(parts) + "\n"


def export_all(proposal: Proposal, resolved: list, cfg, out_dir: Path) -> dict[str, str]:
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    md = compose_markdown(proposal, resolved)
    exported: dict[str, str] = {}

    # Markdown
    md_path = out_dir / "proposal.md"
    md_path.write_text(md, encoding="utf-8")
    exported["markdown"] = "proposal.md"

    if "html" in cfg.export_formats:
        exported["html"] = _export_html(md, out_dir)
    if "docx" in cfg.export_formats:
        exported["docx"] = _export_docx(md, out_dir)
    if "pdf" in cfg.export_formats:
        exported["pdf"] = _export_pdf(md, out_dir)
    return exported


# ---------------- HTML ----------------
def _export_html(md: str, out_dir: Path) -> str:
    try:
        import markdown as mdlib
    except Exception:
        return ""
    body = mdlib.markdown(md, extensions=["tables"])
    html = f"""<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="utf-8">
<title>方案书</title>
<style>
body{{font-family:'Microsoft YaHei',sans-serif;max-width:860px;margin:24px auto;padding:0 16px;line-height:1.7;color:#222}}
img{{max-width:100%;border:1px solid #ddd;border-radius:6px}}
h1{{border-bottom:2px solid #2c5f8a;padding-bottom:8px}}
h2{{color:#2c5f8a;margin-top:1.4em}}
blockquote{{border-left:4px solid #2c5f8a;margin:0;padding-left:12px;color:#555}}
li{{margin:3px 0}}
</style></head><body>
{body}
</body></html>"""
    path = out_dir / "proposal.html"
    path.write_text(html, encoding="utf-8")
    return "proposal.html"


# ---------------- Word ----------------
def _export_docx(md: str, out_dir: Path) -> str:
    try:
        import docx
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except Exception:
        return ""
    d = docx.Document()
    d.styles["Normal"].font.name = "Microsoft YaHei"
    d.styles["Normal"].font.size = Pt(10.5)

    def add_runs(p, text: str):
        pos = 0
        for m in BOLD_RE.finditer(text):
            if m.start() > pos:
                p.add_run(text[pos:m.start()])
            r = p.add_run(m.group(1))
            r.bold = True
            pos = m.end()
        if pos < len(text):
            p.add_run(text[pos:])

    for raw in md.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        m = IMG_RE.match(line)
        if m:
            rel = m.group(2)
            img = out_dir / rel
            if img.exists():
                p = d.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(img), width=docx.shared.Inches(5.5))
            continue
        if line.startswith("## "):
            d.add_heading(line[3:], level=1)
        elif line.startswith("# "):
            d.add_heading(line[2:], level=0)
        elif line.startswith("- "):
            add_runs(d.add_paragraph(style="List Bullet"), line[2:])
        elif line.startswith("> "):
            p = d.add_paragraph()
            r = add_runs(p, line[2:])
            p.runs[-1].italic = True if p.runs else None
        elif line.startswith("*") and line.endswith("*") and len(line) > 2:
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(p, line.strip("*"))
            for r in p.runs:
                r.italic = True
        else:
            add_runs(d.add_paragraph(), line)

    path = out_dir / "proposal.docx"
    d.save(str(path))
    return "proposal.docx"


# ---------------- PDF（matplotlib 兜底，零额外依赖） ----------------
def _export_pdf(md: str, out_dir: Path) -> str:
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from matplotlib.backends.backend_pdf import PdfPages
    except Exception:
        return ""
    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False

    path = out_dir / "proposal.pdf"

    # 分块：文本行 / 图片
    blocks: list[tuple[str, Any]] = []
    for raw in md.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        m = IMG_RE.match(line)
        if m:
            img = out_dir / m.group(2)
            if img.exists():
                blocks.append(("img", img))
            continue
        blocks.append(("text", line))

    A4_W, A4_H = 8.27, 11.69
    with PdfPages(str(path)) as pdf:
        fig = plt.figure(figsize=(A4_W, A4_H))
        y = 0.95
        first = True

        def new_page():
            nonlocal fig, y
            if not first:
                pass
            pdf.savefig(fig)
            plt.close(fig)
            fig = plt.figure(figsize=(A4_W, A4_H))
            y = 0.95

        for kind, payload in blocks:
            if kind == "img":
                if y < 0.42:
                    new_page()
                ax = fig.add_axes([0.12, y - 0.36, 0.76, 0.36])
                ax.imshow(plt.imread(str(payload)))
                ax.axis("off")
                y -= 0.40
            else:
                text = payload
                level = 0
                fontsize = 10.5
                color = "#222222"
                weight = "normal"
                if text.startswith("## "):
                    text = text[3:]; level = 1; fontsize = 14; color = "#2c5f8a"; weight = "bold"
                elif text.startswith("# "):
                    text = text[2:]; level = 0; fontsize = 18; color = "#1c3d5a"; weight = "bold"
                elif text.startswith("- "):
                    text = "• " + text[2:]
                elif text.startswith("> "):
                    text = text[2:]; color = "#555555"
                for wrapped in _wrap_cjk(BOLD_RE.sub(r"\1", text), 38):
                    if y < 0.05:
                        new_page()
                    fig.text(0.06, y, wrapped, fontsize=fontsize, color=color,
                             fontweight=weight, fontfamily="sans-serif")
                    y -= 0.024
                y -= 0.008

        pdf.savefig(fig)
        plt.close(fig)

    return "proposal.pdf"


def _wrap_cjk(text: str, limit: int) -> list[str]:
    """按 CJK 字符宽度折行。"""
    lines: list[str] = []
    cur = ""
    width = 0
    for ch in text:
        w = 2 if ord(ch) > 0x2E80 else 1
        if width + w > limit * 2 and cur:
            lines.append(cur)
            cur, width = "", 0
        cur += ch
        width += w
    if cur:
        lines.append(cur)
    return lines or [text]