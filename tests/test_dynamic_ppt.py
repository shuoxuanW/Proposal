"""动态 PPT 与案例约束回归测试。"""
import pathlib
import sys
import tempfile
import unittest

ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from pptx import Presentation

from src.mv_proposal.agents.writing import WritingAgent
from src.mv_proposal.constraints import check_case_constraints, extract_case_constraints
from src.mv_proposal.models import Chapter, DevicePick, TaskPlan
from src.mv_proposal.ppt_export import export_ppt


def _chapters():
    titles = ["项目概述", "需求分析", "总体方案设计", "系统组成与光路布局",
              "设备选型说明", "系统软件与算法", "安装调试与验收", "服务与售后",
              "检测效果与缺陷样张", "验证数据与验收指标", "风险与注意事项", "设备参数汇总"]
    return [Chapter(id=f"ch{i+1}", title=t,
                    content=f"## {t}\n\n- 本页内容 {i+1}\n- 156×94mm / 0.008mm/pixel / 130±20mm")
            for i, t in enumerate(titles)]


def _plan():
    return TaskPlan(
        project_name="动态PPT测试方案",
        requirement_summary="金属加工件 156×94mm，检测牙套漏装，精度0.008mm/pixel，三工位",
        detection_target="牙套漏装/残留",
        key_metrics={"精度": "0.008mm/pixel", "速度": "30件/分钟", "视场": "40×27mm", "工作距": "130±20mm"},
        devices=[
            DevicePick(device_type="camera", label="相机", spec_hint="MV-CS200-10GM", quantity=3),
            DevicePick(device_type="lens", label="镜头", spec_hint="YD-1250-25MP 50mm FA", quantity=3),
            DevicePick(device_type="light_source", label="光源", spec_hint="白色圆顶光源", quantity=2),
            DevicePick(device_type="controller", label="控制器", spec_hint="四通道光源控制器", quantity=1),
            DevicePick(device_type="ipc", label="工控机", spec_hint="I7/16G/512G", quantity=1),
        ],
        chapters=_chapters(),
        customer_name="测试客户",
        product_name="金属加工件",
        product_size="156×94mm",
        requirements=["4处牙套孔漏装", "牙套残留小尾巴", "2处小孔毛边"],
        case_constraints={
            "product_size": "156×94mm", "accuracy": "0.008mm/pixel",
            "cycle_time": "30件/分钟", "fov": "40×27mm", "working_distance": "130±20mm",
            "camera_count": 3, "camera_model": "MV-CS200-10GM",
            "lens": "YD-1250-25MP 50mm FA", "light_type": "白色圆顶光源",
            "controller": "四通道光源控制器", "ipc": "I7/16G/512G",
            "station_count": 3, "source": "命中案例",
        },
    )


class ConstraintsTest(unittest.TestCase):
    def test_extract_case_constraints(self):
        case = ("产品名称：金属加工件\n产品尺寸：156×94mm\n检测精度：0.008mm/pixel\n"
                "节拍：30件/分钟\n视野：40×27mm\n工作距离：130±20mm\n"
                "相机：3台\n镜头：YD-1250-25MP 50mm FA\n光源：白色圆顶光源\n"
                "控制器：四通道光源控制器\n工控机：IPC-I7T10\n工位1、工位2、工位3")
        plan = _plan()
        cc = extract_case_constraints(case, plan)
        self.assertIn("156", cc.product_size)
        self.assertIn("0.008", cc.accuracy)
        self.assertIn("30", cc.cycle_time)
        self.assertIn("130", cc.working_distance)
        self.assertEqual(cc.station_count, 3)
        self.assertIn("圆顶", cc.light_type)

    def test_camera_wd_preferred_over_light_distance(self):
        plan = TaskPlan(
            project_name="FPC",
            requirement_summary="光源到物料距离30±10mm",
            key_metrics={"工作距": "30±10mm"},
        )
        cc = extract_case_constraints("相机工作距离373±10mm，光源距离被测物30±10mm", plan)
        self.assertEqual(cc.working_distance, "373±10mm")

    def test_light_spec_types_override_generic_light_type(self):
        plan = _plan()
        plan.case_constraints["light_type"] = "环形"
        plan.case_constraints["light_specs"] = [{"type": "条形"}, {"type": "同轴"}]
        chapters = [Chapter(id="ch1", title="项目概述", content="条形光源和同轴光源")]
        issues = check_case_constraints(plan, chapters)
        self.assertFalse(any(i.code == "OPTICAL_MISMATCH" for i in issues))

    def test_descriptive_device_hints_are_not_models(self):
        plan = TaskPlan(
            project_name="描述性设备测试",
            requirement_summary="CST 工业相机竖直安装，配套镜头垂直向下，控制器用于输出剔除信号",
            devices=[
                DevicePick(device_type="camera", label="工业相机", spec_hint="CST 工业相机，竖直安装", quantity=1),
                DevicePick(device_type="lens", label="镜头", spec_hint="配套镜头，垂直向下", quantity=1),
                DevicePick(device_type="controller", label="控制器", spec_hint="用于输出剔除信号", quantity=1),
            ],
        )
        cc = extract_case_constraints("CST 工业相机竖直安装，配套镜头垂直向下，控制器用于输出剔除信号", plan)
        self.assertEqual(cc.camera_model, "")
        self.assertEqual(cc.lens, "")
        self.assertEqual(cc.controller, "")
        plan.case_constraints = cc.to_dict()
        issues = check_case_constraints(plan, [])
        # 不应因为描述性文字而产生相机型号/镜头/控制器缺失告警
        self.assertFalse(any(i.code == "PARAM_MISSING" and "相机型号" in i.message for i in issues))

    def test_no_coverage_for_plain_requirement(self):
        plan = TaskPlan(case_constraints={"source": "需求文本", "cycle_time": "5个/秒"})
        self.assertEqual(check_case_constraints(plan, []), [])

    def test_coverage_reports_missing(self):
        plan = _plan()
        chapters = [Chapter(id="ch1", title="项目概述", content="产品尺寸 156×94mm，精度 0.008mm/pixel")]
        issues = check_case_constraints(plan, chapters)
        self.assertTrue(any(i.code == "PARAM_MISSING" for i in issues))
        self.assertTrue(any(i.code == "PARAM_COVERAGE" and "%" in i.message for i in issues))


class DeviceSummaryTest(unittest.TestCase):
    def test_device_summary_fills_missing_models(self):
        plan = _plan()
        ch = Chapter(id="ch12", title="设备参数汇总", content="## 设备参数汇总\n\n无")
        out = WritingAgent(None, None)._ensure_device_summary(plan, [ch])
        self.assertIn("MV-CS200-10GM", out[0].content)
        self.assertIn("YD-1250-25MP 50mm FA", out[0].content)
        self.assertIn("130±20mm", out[0].content)


class MetricsPriorityTest(unittest.TestCase):
    def test_camera_wd_wins_over_light_distance(self):
        from src.mv_proposal.ppt.planner import _metrics_table
        plan = _plan()
        plan.case_constraints["working_distance"] = "373±10mm"
        plan.key_metrics["工作距"] = "30±10mm"
        table = _metrics_table(plan)
        row = next(r for r in table if r[0] == "相机工作距离")
        self.assertEqual(row[1], "373±10mm")
        text = WritingAgent._validation_section(plan)
        self.assertIn("373±10mm", text)


class DynamicPPTTest(unittest.TestCase):
    def test_dynamic_ppt_has_appendix_pages(self):
        plan = _plan()
        with tempfile.TemporaryDirectory() as td:
            export_ppt(plan, [], ROOT / "data" / "templates" / "视觉方案书格式.pptx",
                       pathlib.Path(td), chapters=plan.chapters)
            prs = Presentation(str(pathlib.Path(td) / "proposal.pptx"))
            self.assertGreaterEqual(len(prs.slides), 12)
            texts = []
            for slide in prs.slides:
                for sh in slide.shapes:
                    if sh.has_text_frame:
                        texts.append(sh.text_frame.text)
                    if sh.has_table:
                        for row in sh.table.rows:
                            texts.append(" | ".join(c.text for c in row.cells))
            text = "\n".join(texts)
            for title in ("检测效果与缺陷样张", "验证数据与验收指标", "风险与注意事项", "设备参数汇总"):
                self.assertIn(title, text)
            self.assertIn("检测精度", text)
            self.assertIn("MV-CS200-10GM", text)


if __name__ == "__main__":
    unittest.main()
