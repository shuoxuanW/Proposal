"""端到端冒烟测试（mock 模式，不调用外部 API）。"""
import pathlib
import sys
import tempfile
import unittest

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))

from src.mv_proposal.config import Config
from src.mv_proposal.pipeline import generate


class PipelineTest(unittest.TestCase):
    def setUp(self):
        self.cfg = Config.load()
        sample = self.cfg.samples_dir / "requirement_example.md"
        self.text = sample.read_text(encoding="utf-8") if sample.exists() else "检测金属表面缺陷，精度0.05mm，速度120件/分钟"

    def test_mock_pipeline_end_to_end(self):
        with tempfile.TemporaryDirectory() as td:
            res = generate(self.text, cfg=self.cfg, llm_provider="mock", out_dir=pathlib.Path(td))
            self.assertTrue(res["report"].passed)
            for name in ("proposal.md", "proposal.html", "proposal.docx", "proposal.pdf",
                         "linkage.json", "manifest.json", "proposal.json"):
                self.assertTrue((res["run_dir"] / name).exists(), name)
            self.assertTrue((res["run_dir"] / "figures" / "optic-1.png").exists())
            # 动态章节 + 动态 PPT 附录页
            self.assertGreaterEqual(len(res["plan"].chapters), 12)
            from pptx import Presentation
            prs = Presentation(str(res["run_dir"] / "proposal.pptx"))
            self.assertGreaterEqual(len(prs.slides), 12)
            ppt_text = "\n".join(
                sh.text_frame.text for slide in prs.slides for sh in slide.shapes
                if sh.has_text_frame)
            for title in ("检测效果与缺陷样张", "验证数据与验收指标", "风险与注意事项", "设备参数汇总"):
                self.assertIn(title, ppt_text)
            # 图文联动注册表应包含图纸映射
            import json
            link = json.loads((res["run_dir"] / "linkage.json").read_text(encoding="utf-8"))
            self.assertTrue(any(e["kind"] == "fig" for e in link["entries"]))
            # 案例召回结果必须同时出现在返回值和落盘 manifest.json 中
            self.assertIsInstance(res["matched_cases"], list)
            for c in res["matched_cases"]:
                self.assertIn("source", c)
                self.assertIn("score", c)
            disk_manifest = json.loads((res["run_dir"] / "manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(disk_manifest.get("matched_cases"), res["matched_cases"])


class FeedbackTest(unittest.TestCase):
    def test_ingest_text(self):
        import pathlib, tempfile
        from src.mv_proposal.feedback import FeedbackManager
        root = pathlib.Path(tempfile.mkdtemp())
        cfg = Config(project_root=root)
        cfg.ensure_dirs()
        fm = FeedbackManager(cfg)
        r = fm.ingest_text("人工补充：高反光表面建议采用同轴光照明", source="人工修改.md", collection="devices")
        self.assertEqual(r["collection"], "devices")
        self.assertGreater(fm.ctx.kb.count(), 0)


if __name__ == "__main__":
    unittest.main()